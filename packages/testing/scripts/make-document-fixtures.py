"""Makes the Office fixtures of the SET 10 document tests with independent
producers (python-docx, python-pptx, openpyxl), so that Jupiter's readers are
tested on files it did not write. Run from the repository root:

    python3 packages/testing/scripts/make-document-fixtures.py

The PDF and the PNG come from headless Chromium: make-document-fixtures.mjs.
"""
import datetime
import os

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

HERE = os.path.join(os.path.dirname(__file__), '..', 'fixtures', 'documents')

doc = Document()
doc.core_properties.title = 'Mission Briefing'
doc.core_properties.author = 'Fixture Author'
doc.add_heading('Mission Briefing', level=1)
doc.add_paragraph('Jupiter has 95 known moons.')
doc.add_heading('Moons', level=2)
doc.add_paragraph('Io is volcanic.', style='List Bullet')
doc.add_paragraph('Europa has an ice shell.', style='List Bullet')
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = 'Moon'
table.cell(0, 1).text = 'Diameter (km)'
table.cell(1, 0).text = 'Ganymede'
table.cell(1, 1).text = '5268'
doc.add_paragraph('ภาษาไทย: ดาวพฤหัสบดี')
doc.save(os.path.join(HERE, 'briefing.docx'))

deck = Presentation()
deck.core_properties.title = 'Quarterly Review'
deck.core_properties.author = 'Fixture Author'
title = deck.slides.add_slide(deck.slide_layouts[0])
title.shapes.title.text = 'Quarterly Review'
title.placeholders[1].text = 'Second quarter'
title.notes_slide.notes_text_frame.text = 'Welcome everyone.'
content = deck.slides.add_slide(deck.slide_layouts[1])
content.shapes.title.text = 'Results'
body = content.placeholders[1].text_frame
body.text = 'Revenue grew 12%'
body.add_paragraph().text = 'Costs fell 3%'
content.notes_slide.notes_text_frame.text = 'Mention the new office.'
content.shapes.add_picture(os.path.join(HERE, 'chart.png'), Inches(6), Inches(2), width=Inches(3))
deck.save(os.path.join(HERE, 'review.pptx'))

book = Workbook()
sheet = book.active
sheet.title = 'Budget'
sheet.append(['Item', 'Amount', 'Date'])
sheet.append(['Telescope', 1200, datetime.date(2026, 3, 1)])
sheet.append(['Filters', 300.5, datetime.date(2026, 3, 2)])
sheet.append(['Total', '=SUM(B2:B3)', None])
book.properties.title = 'Budget 2026'
book.create_sheet('Notes').append(['Approved by the team'])
book.save(os.path.join(HERE, 'budget.xlsx'))
print('ok')
