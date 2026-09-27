"""Opens an Office file with an independent, standards-based parser
(python-docx, python-pptx or openpyxl) and prints what it found as JSON.
Used by the SET 10 tests to prove that documents Jupiter writes open in a
program other than Jupiter.

    python3 check-office.py <file> <docx|pptx|xlsx>
"""
import datetime
import json
import sys


def docx_summary(path):
    from docx import Document

    doc = Document(path)
    return {
        'title': doc.core_properties.title,
        'paragraphs': [{'style': p.style.name, 'text': p.text} for p in doc.paragraphs],
        'tables': [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables],
    }


def pptx_summary(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(path)
    slides = []
    for slide in deck.slides:
        texts = []
        for shape in slide.placeholders:
            if shape.placeholder_format.idx != 0 and shape.has_text_frame:
                texts.append(shape.text_frame.text)
        slides.append(
            {
                'layout': slide.slide_layout.name,
                'title': slide.shapes.title.text if slide.shapes.title is not None else None,
                'texts': texts,
                'notes': slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else None,
                'pictures': sum(1 for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE),
            }
        )
    return {
        'title': deck.core_properties.title,
        'width': deck.slide_width,
        'height': deck.slide_height,
        'slides': slides,
    }


def xlsx_summary(path):
    from openpyxl import load_workbook

    book = load_workbook(path)
    sheets = []
    for sheet in book.worksheets:
        rows = []
        for row in sheet.iter_rows():
            values = []
            for cell in row:
                value = cell.value
                if isinstance(value, (datetime.datetime, datetime.date)):
                    value = value.isoformat()
                values.append({'value': value, 'type': cell.data_type, 'format': cell.number_format})
            rows.append(values)
        sheets.append({'name': sheet.title, 'rows': rows})
    return {'title': book.properties.title, 'sheets': sheets}


if __name__ == '__main__':
    target, kind = sys.argv[1], sys.argv[2]
    summary = {'docx': docx_summary, 'pptx': pptx_summary, 'xlsx': xlsx_summary}[kind](target)
    print(json.dumps(summary, ensure_ascii=False))
